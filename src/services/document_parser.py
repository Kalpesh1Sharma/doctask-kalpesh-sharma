"""Restricted, truthful parsers for the Phase 6A source formats."""

from __future__ import annotations

import io
import zipfile
from dataclasses import dataclass
from pathlib import Path
from xml.etree import ElementTree

_WORD = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
_DRAWING = "http://schemas.openxmlformats.org/drawingml/2006/main"
_OFFICE = "urn:schemas-microsoft-com:office:office"
_NS = {"w": _WORD}
_TXT_READ_CHARS = 1024 * 1024


class DocumentParseError(RuntimeError):
    pass


class UnsupportedDocumentTypeError(DocumentParseError):
    pass


class UnsupportedDocxContentError(DocumentParseError):
    pass


class ScannedPdfError(DocumentParseError):
    pass


@dataclass(frozen=True)
class ParsedSegment:
    text: str
    page_number: int | None


@dataclass(frozen=True)
class ParsedDocument:
    mime_type: str
    segments: tuple[ParsedSegment, ...]

    @property
    def text(self) -> str:
        if all(segment.page_number is None for segment in self.segments):
            return "".join(segment.text for segment in self.segments)
        return "\n".join(segment.text for segment in self.segments)


def detect_source_type(data: bytes) -> str:
    """Detect exactly the supported types from bytes, never a filename."""

    if data.startswith(b"%PDF-"):
        return "application/pdf"
    if data.startswith((b"PK\x03\x04", b"PK\x05\x06")):
        try:
            with zipfile.ZipFile(io.BytesIO(data)) as archive:
                names = set(archive.namelist())
                if "[Content_Types].xml" in names and "word/document.xml" in names:
                    return "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        except zipfile.BadZipFile as error:
            raise UnsupportedDocumentTypeError("ZIP source is not a valid DOCX") from error
    try:
        decoded = data.decode("utf-8-sig", "strict")
    except UnicodeDecodeError as error:
        raise UnsupportedDocumentTypeError("source is not PDF, DOCX, or strict UTF-8 TXT") from error
    if decoded.startswith("\ufeff"):
        # utf-8-sig removes only one leading BOM.  A remaining U+FEFF is text.
        raise UnsupportedDocumentTypeError("TXT accepts only an optional UTF-8 byte-order mark")
    return "text/plain; charset=utf-8"


def detect_source_type_path(path: Path) -> str:
    """Detect a stored source without materializing the whole file as bytes."""

    source = Path(path)
    try:
        with source.open("rb") as stream:
            signature = stream.read(5)
    except OSError as error:
        raise DocumentParseError("source cannot be opened") from error

    if signature.startswith(b"%PDF-"):
        return "application/pdf"
    if signature.startswith((b"PK\x03\x04", b"PK\x05\x06")):
        try:
            with zipfile.ZipFile(source) as archive:
                names = set(archive.namelist())
                if "[Content_Types].xml" in names and "word/document.xml" in names:
                    return "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        except zipfile.BadZipFile as error:
            raise UnsupportedDocumentTypeError("ZIP source is not a valid DOCX") from error

    # TXT validation is performed by the path parser itself so the file is read
    # through a text decoder rather than duplicated as one large bytes object.
    return "text/plain; charset=utf-8"


class DocumentParser:
    """Parse the limited supported set while rejecting omitted DOCX content."""

    def parse_path(self, path: Path) -> ParsedDocument:
        source = Path(path)
        mime_type = detect_source_type_path(source)
        if mime_type == "application/pdf":
            return self._parse_pdf_path(source)
        if mime_type.startswith("application/vnd.openxmlformats"):
            return self._parse_docx_path(source)
        return self._parse_txt_path(source)

    def parse_bytes(self, data: bytes) -> ParsedDocument:
        mime_type = detect_source_type(data)
        if mime_type == "application/pdf":
            return self._parse_pdf(data)
        if mime_type.startswith("application/vnd.openxmlformats"):
            return self._parse_docx(data)
        return self._parse_txt(data)

    @staticmethod
    def _parse_txt(data: bytes) -> ParsedDocument:
        try:
            text = data.decode("utf-8-sig", "strict")
        except UnicodeDecodeError as error:
            raise UnsupportedDocumentTypeError("TXT source must be strict UTF-8") from error
        return ParsedDocument("text/plain; charset=utf-8", (ParsedSegment(text, None),))

    @staticmethod
    def _parse_txt_path(path: Path) -> ParsedDocument:
        try:
            with path.open("r", encoding="utf-8-sig", errors="strict", newline="") as stream:
                segments: list[ParsedSegment] = []
                while text := stream.read(_TXT_READ_CHARS):
                    if not segments and text.startswith("\ufeff"):
                        raise UnsupportedDocumentTypeError(
                            "TXT accepts only an optional UTF-8 byte-order mark",
                        )
                    segments.append(ParsedSegment(text, None))
        except UnicodeDecodeError as error:
            raise UnsupportedDocumentTypeError("TXT source must be strict UTF-8") from error
        except OSError as error:
            raise DocumentParseError("TXT source cannot be read") from error
        return ParsedDocument("text/plain; charset=utf-8", tuple(segments))

    @staticmethod
    def _parse_pdf(data: bytes) -> ParsedDocument:
        try:
            from pypdf import PdfReader
        except ImportError as error:  # pragma: no cover - packaging assertion
            raise DocumentParseError("pypdf must be installed for PDF ingestion") from error
        try:
            reader = PdfReader(io.BytesIO(data))
            segments = tuple(
                ParsedSegment(page.extract_text() or "", index + 1)
                for index, page in enumerate(reader.pages)
            )
        except Exception as error:  # pypdf exposes several parse-specific exceptions
            raise DocumentParseError("PDF cannot be parsed") from error
        if not any(segment.text.strip() for segment in segments):
            raise ScannedPdfError("PDF contains no extractable text; OCR is not supported")
        return ParsedDocument("application/pdf", segments)

    @staticmethod
    def _parse_pdf_path(path: Path) -> ParsedDocument:
        try:
            from pypdf import PdfReader
        except ImportError as error:  # pragma: no cover - packaging assertion
            raise DocumentParseError("pypdf must be installed for PDF ingestion") from error
        try:
            reader = PdfReader(str(path))
            segments = tuple(
                ParsedSegment(page.extract_text() or "", index + 1)
                for index, page in enumerate(reader.pages)
            )
        except Exception as error:  # pypdf exposes several parse-specific exceptions
            raise DocumentParseError("PDF cannot be parsed") from error
        if not any(segment.text.strip() for segment in segments):
            raise ScannedPdfError("PDF contains no extractable text; OCR is not supported")
        return ParsedDocument("application/pdf", segments)

    def _parse_docx(self, data: bytes) -> ParsedDocument:
        # Keep python-docx as an explicit dependency / supported format parser,
        # while OOXML inspection below supplies the fail-closed coverage rule.
        try:
            import docx  # noqa: F401
        except ImportError as error:  # pragma: no cover - packaging assertion
            raise DocumentParseError("python-docx must be installed for DOCX ingestion") from error
        try:
            with zipfile.ZipFile(io.BytesIO(data)) as archive:
                names = set(archive.namelist())
                if "word/document.xml" not in names:
                    raise UnsupportedDocumentTypeError("DOCX is missing word/document.xml")
                self._reject_unsupported_docx_parts(archive, names)
                root = ElementTree.fromstring(archive.read("word/document.xml"))
        except zipfile.BadZipFile as error:
            raise UnsupportedDocumentTypeError("DOCX archive is malformed") from error
        except ElementTree.ParseError as error:
            raise DocumentParseError("DOCX XML is malformed") from error
        self._reject_unsupported_docx_markup(root)
        body = root.find(f"{{{_WORD}}}body")
        if body is None:
            raise DocumentParseError("DOCX has no main body")
        blocks: list[str] = []
        for child in body:
            if child.tag == f"{{{_WORD}}}p":
                blocks.append(self._paragraph_text(child))
            elif child.tag == f"{{{_WORD}}}tbl":
                blocks.extend(self._table_blocks(child))
            elif child.tag != f"{{{_WORD}}}sectPr":
                raise UnsupportedDocxContentError("DOCX body contains an unsupported construct")
        text = "\n".join(block for block in blocks if block != "")
        return ParsedDocument(
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            (ParsedSegment(text, None),),
        )

    def _parse_docx_path(self, path: Path) -> ParsedDocument:
        try:
            import docx  # noqa: F401
        except ImportError as error:  # pragma: no cover - packaging assertion
            raise DocumentParseError("python-docx must be installed for DOCX ingestion") from error
        try:
            with zipfile.ZipFile(path) as archive:
                names = set(archive.namelist())
                if "word/document.xml" not in names:
                    raise UnsupportedDocumentTypeError("DOCX is missing word/document.xml")
                self._reject_unsupported_docx_parts(archive, names)
                root = ElementTree.fromstring(archive.read("word/document.xml"))
        except zipfile.BadZipFile as error:
            raise UnsupportedDocumentTypeError("DOCX archive is malformed") from error
        except ElementTree.ParseError as error:
            raise DocumentParseError("DOCX XML is malformed") from error
        self._reject_unsupported_docx_markup(root)
        body = root.find(f"{{{_WORD}}}body")
        if body is None:
            raise DocumentParseError("DOCX has no main body")
        blocks: list[str] = []
        for child in body:
            if child.tag == f"{{{_WORD}}}p":
                blocks.append(self._paragraph_text(child))
            elif child.tag == f"{{{_WORD}}}tbl":
                blocks.extend(self._table_blocks(child))
            elif child.tag != f"{{{_WORD}}}sectPr":
                raise UnsupportedDocxContentError("DOCX body contains an unsupported construct")
        text = "\n".join(block for block in blocks if block != "")
        return ParsedDocument(
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            (ParsedSegment(text, None),),
        )

    @staticmethod
    def _part_has_text(data: bytes) -> bool:
        try:
            root = ElementTree.fromstring(data)
        except ElementTree.ParseError as error:
            raise UnsupportedDocxContentError("unsupported DOCX part is malformed") from error
        return any((node.text or "").strip() for node in root.iter(f"{{{_WORD}}}t"))

    def _reject_unsupported_docx_parts(self, archive: zipfile.ZipFile, names: set[str]) -> None:
        text_parts = [
            name
            for name in names
            if name.startswith(("word/header", "word/footer"))
            or name in {"word/footnotes.xml", "word/endnotes.xml", "word/comments.xml"}
        ]
        for name in text_parts:
            if self._part_has_text(archive.read(name)):
                raise UnsupportedDocxContentError(f"DOCX contains unsupported text-bearing {name}")
        if any(
            name.startswith(("word/embeddings/", "word/media/", "word/activeX/"))
            for name in names
        ):
            raise UnsupportedDocxContentError("DOCX contains unsupported embedded or drawing content")

    @staticmethod
    def _reject_unsupported_docx_markup(root: ElementTree.Element) -> None:
        rejected = {
            f"{{{_WORD}}}ins",
            f"{{{_WORD}}}del",
            f"{{{_WORD}}}moveFrom",
            f"{{{_WORD}}}moveTo",
            f"{{{_WORD}}}fldChar",
            f"{{{_WORD}}}instrText",
            f"{{{_WORD}}}object",
            f"{{{_WORD}}}pict",
            f"{{{_WORD}}}txbxContent",
            f"{{{_DRAWING}}}blip",
            f"{{{_OFFICE}}}OLEObject",
        }
        for node in root.iter():
            if node.tag in rejected or node.tag.endswith("}drawing"):
                raise UnsupportedDocxContentError("DOCX contains unsupported tracked, field, or drawing content")

    @staticmethod
    def _paragraph_text(paragraph: ElementTree.Element) -> str:
        # w:t preserves literal body text in run order.  Tabs and explicit line
        # breaks are represented truthfully rather than discarded.
        parts: list[str] = []
        for node in paragraph.iter():
            if node.tag == f"{{{_WORD}}}t":
                parts.append(node.text or "")
            elif node.tag == f"{{{_WORD}}}tab":
                parts.append("\t")
            elif node.tag in {f"{{{_WORD}}}br", f"{{{_WORD}}}cr"}:
                parts.append("\n")
        return "".join(parts)

    def _table_blocks(self, table: ElementTree.Element) -> list[str]:
        blocks: list[str] = []
        for row in table.findall("w:tr", _NS):
            cells: list[str] = []
            for cell in row.findall("w:tc", _NS):
                content: list[str] = []
                for child in cell:
                    if child.tag == f"{{{_WORD}}}p":
                        content.append(self._paragraph_text(child))
                    elif child.tag == f"{{{_WORD}}}tbl":
                        content.extend(self._table_blocks(child))
                    elif child.tag != f"{{{_WORD}}}tcPr":
                        raise UnsupportedDocxContentError("DOCX table cell contains unsupported content")
                cells.append("\n".join(content))
            blocks.append("\t".join(cells))
        return blocks
